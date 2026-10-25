from openpyxl import load_workbook

from planilha import COLUNAS, gerar_planilha

LINHAS = [
    {"numero": "26BR0000012345-6", "data_registro": "14/03/2026", "qtd": 480.0, "peso": 12345.678,
     "ncm": "4011.10.00, 4011.20.90", "cnpj": "12.345.678/0001-90", "nome": "COMERCIAL EXEMPLO LTDA",
     "referencia": "IMP 1099B"},
    {"numero": "26BR0000067890-1", "data_registro": "", "qtd": "", "peso": "",
     "ncm": "", "cnpj": "", "nome": "", "referencia": ""},
]


def test_uma_linha_por_duimp_com_cabecalho(tmp_path):
    arq = tmp_path / "duimp.xlsx"
    gerar_planilha(LINHAS, arq, "12.345.678/0001-90", "01/01/2026", "31/03/2026")

    ws = load_workbook(arq)["DUIMP"]
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0] == tuple(titulo for _, titulo, _, _ in COLUNAS)
    assert linhas[1][0] == "26BR0000012345-6"
    assert linhas[1][3] == 12345.678
    assert len(linhas) == 3
    assert ws["D2"].number_format == "#,##0.000"
    assert ws.auto_filter.ref == "A1:H3"


def test_resumo(tmp_path):
    arq = tmp_path / "duimp.xlsx"
    gerar_planilha(LINHAS, arq, "12.345.678/0001-90", "01/01/2026", "31/03/2026")

    resumo = {l[0]: l[1] for l in load_workbook(arq)["Resumo"].iter_rows(values_only=True)}
    assert resumo["Total de DUIMPs"] == 2
    assert resumo["Período"] == "01/01/2026 a 31/03/2026"
