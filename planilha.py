"""Planilha Excel com uma linha por DUIMP e uma aba de resumo da extracao."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# (campo, titulo, largura, centralizado)
COLUNAS = [
    ("numero", "Número da DUIMP", 22, True),
    ("data_registro", "Data de registro", 14, True),
    ("qtd", "Qtd. unidade comercializada", 18, True),
    ("peso", "Peso bruto total (kg)", 18, True),
    ("ncm", "NCM", 38, False),
    ("cnpj", "CNPJ do adquirente", 22, False),
    ("nome", "Nome do adquirente", 30, False),
    ("referencia", "Referência interna", 18, True),
]

AZUL = "1F4E79"
BORDA = Border(*(Side(style="thin", color="CCCCCC"),) * 4)


def gerar_planilha(linhas, caminho, cnpj, data_ini, data_fim):
    wb = Workbook()
    ws = wb.active
    ws.title = "DUIMP"

    for col, (_, titulo, largura, _) in enumerate(COLUNAS, 1):
        cel = ws.cell(row=1, column=col, value=titulo)
        cel.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cel.fill = PatternFill("solid", start_color=AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDA
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.row_dimensions[1].height = 36

    zebra = PatternFill("solid", start_color="EBF3FB")
    for lin, dados in enumerate(linhas, 2):
        for col, (campo, _, _, centro) in enumerate(COLUNAS, 1):
            valor = dados.get(campo, "")
            cel = ws.cell(row=lin, column=col, value=valor)
            cel.font = Font(name="Arial", size=10)
            cel.border = BORDA
            cel.alignment = Alignment(horizontal="center" if centro else "left", vertical="center",
                                      wrap_text=not centro)
            if lin % 2 == 0:
                cel.fill = zebra
            if campo == "qtd" and isinstance(valor, (int, float)):
                cel.number_format = "#,##0"
            if campo == "peso" and isinstance(valor, float):
                cel.number_format = "#,##0.000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{len(linhas) + 1}"

    resumo = wb.create_sheet("Resumo")
    resumo.append(["Relatório DUIMP"])
    resumo["A1"].font = Font(name="Arial", bold=True, size=14, color=AZUL)
    for rotulo, valor in [
        ("CNPJ importador", cnpj),
        ("Período", f"{data_ini} a {data_fim}"),
        ("Total de DUIMPs", len(linhas)),
        ("Data da extração", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]:
        resumo.append([rotulo, valor])
    resumo.column_dimensions["A"].width = 22
    resumo.column_dimensions["B"].width = 40
    wb.save(caminho)
